"""阶段一：解析。"""
from __future__ import annotations

import docx
from docx.enum.style import WD_STYLE_TYPE

from main import clean_text, docx_parser, load_directory, markdown_parser, text_parser


def test_clean_text_unifies_newlines_in_the_right_order():
    # 先换 \r\n 再换 \r：顺序反了的话，一个换行会变成两个
    assert clean_text("a\r\n\r\nb\rc\r\n  ") == "a\n\nb\nc"


def test_text_parser_reads_utf8(tmp_path):
    path = tmp_path / "a.txt"
    path.write_text("你好\n世界", encoding="utf-8")
    doc = text_parser(path)
    assert doc.content == "你好\n世界"
    assert doc.meta_data["encoding"] == "utf-8"


def test_text_parser_falls_back_to_gb18030(tmp_path):
    path = tmp_path / "会议记录.txt"
    path.write_bytes("2026-10-20 周会\r\n结论：先做本地文件".encode("gbk"))
    doc = text_parser(path)
    assert "周会" in doc.content
    assert doc.meta_data["encoding"] == "gb18030"


def test_markdown_parser_keeps_structure(tmp_path):
    path = tmp_path / "note.md"
    path.write_text("# 标题\n\n正文", encoding="utf-8")
    doc = markdown_parser(path)
    assert doc.content.startswith("# 标题")
    assert doc.meta_data["parser"] == "markdown"


def test_docx_parser_turns_headings_and_code_into_markdown(tmp_path):
    d = docx.Document()
    d.styles.add_style("HTML Preformatted", WD_STYLE_TYPE.PARAGRAPH)
    d.add_paragraph("题目 1", style="Heading 3")
    d.add_paragraph("题目描述")
    d.add_paragraph("cpp")                                   # 语言标签
    d.add_paragraph("int main() {", style="HTML Preformatted")
    d.add_paragraph("    return 0;}", style="HTML Preformatted")
    d.add_paragraph("代码解释")
    path = tmp_path / "题目.docx"
    d.save(path)

    content = docx_parser(path).content
    assert "### 题目 1" in content
    assert "```cpp\nint main() {\n    return 0;}\n```" in content   # 缩进保留
    assert content.index("```cpp") < content.index("代码解释")


def test_load_directory_sorts_files_into_three_buckets(tmp_path):
    (tmp_path / "ok.md").write_text("# 有内容", encoding="utf-8")
    (tmp_path / "blank.md").write_text("  \n\n ", encoding="utf-8")   # 清洗后为空
    (tmp_path / "data.csv").write_text("a,b", encoding="utf-8")       # 不支持的格式

    documents, failed, empty = load_directory(str(tmp_path))

    assert [d.name for d in documents] == ["ok.md"]
    assert [d.name for d in empty] == ["blank.md"]
    assert "不支持" in failed[str(tmp_path / "data.csv")]
