from dataclasses import field, dataclass
from pathlib import Path
import hashlib

import docx  # python-docx 库。注意不要写 from docx import Document，会和我们自己的 Document 重名
from ocr_service import recognize_file



@dataclass
class Document:
    id:str
    name:str
    content:str
    source_path:str
    meta_data:dict = field(default_factory=dict)

def _ocr_document(path, provider=None, mode="LIVE"):
    result = recognize_file(path, provider=provider, mode=mode)
    return Document(
        id=md5_calculate_doc_id(str(path)),
        name=path.name,
        content=result["text"],
        source_path=str(path),
        meta_data={
            "parser": "ocr",
            "ocr_provider": result["ocr_provider"],
            "page_count": result["page_count"],
            "ocr_result_path": result["result_path"],
        },
    )


def pdf_parser(path, provider=None, mode="LIVE")->Document:
    return _ocr_document(path, provider=provider, mode=mode)

def picture_parser(path, provider=None, mode="LIVE")->Document:
    return _ocr_document(path, provider=provider, mode=mode)

# 哪些段落样式算"代码"。以后遇到别的写法，往这里加一个名字就行
CODE_STYLES = {"HTML Preformatted"}

# 代码块前面单独一行的语言标签（从网页复制代码时常会带进来）
CODE_LANGS = {"c", "cpp", "c++", "java", "python", "py", "go", "sql",
              "bash", "shell", "sh", "javascript", "js", "json"}


def docx_parser(path) -> Document:
    """docx → Document：标题样式转成 #，代码样式转成 ``` 代码块。"""
    d = docx.Document(str(path))
    blocks = []        # 一块一块的 Markdown，最后用空行连起来
    code_lines = []    # 正在收集的代码行；不为空 = 当前正在代码块里
    code_lang = ""     # 当前代码块的语言，比如 "cpp"

    def flush_code():
        """把收集好的代码行包成 ```代码块```，放进 blocks，然后清空。"""
        nonlocal code_lines, code_lang   # 说明这里改的是外层函数的变量
        if code_lines:
            blocks.append(f"```{code_lang}\n" + "\n".join(code_lines) + "\n```")
            code_lines, code_lang = [], ""

    for p in d.paragraphs:
        style = p.style.name if p.style is not None else ""

        # ① 代码段落：先攒着，不急着输出
        if style in CODE_STYLES:
            # 刚进入代码块，且上一块正好是 "cpp" 这种标签 → 拿来当语言
            if not code_lines and blocks and blocks[-1].lower() in CODE_LANGS:
                code_lang = blocks.pop().lower()
            code_lines.append(p.text)    # 代码原样保留，缩进和空行都有意义
            continue

        # ② 走到这里说明不是代码了，如果刚才在代码块里，先收尾
        flush_code()

        text = p.text.strip()
        if not text:
            continue                     # 正文里的空段落跳过，块之间本来就会加空行

        # ③ 标题："Heading 3" → "### 文字"
        parts = style.split()            # "Heading 3" → ["Heading", "3"]
        if len(parts) == 2 and parts[0] == "Heading" and parts[1].isdigit():
            blocks.append("#" * int(parts[1]) + " " + text)
        else:
            blocks.append(text)          # ④ 普通段落原样保留

    flush_code()                         # 文档以代码结尾时，最后一块也要收尾

    return Document(
        id=md5_calculate_doc_id(str(path)),
        name=path.name,
        content="\n\n".join(blocks),
        source_path=str(path),
        meta_data={"parser": "docx"},
    )

def read_text_file(path)->tuple[str,str]:
    data = path.read_bytes()
    try:
        text = data.decode("utf-8")
        encoding = "utf-8"
    except UnicodeDecodeError:
        text = data.decode("gb18030")
        encoding = "gb18030"
    return text,encoding


def text_parser(path) ->Document:
    text, encoding = read_text_file(path)
    return Document(
        id=md5_calculate_doc_id(str(path)),
        name=path.name,
        content=text,
        source_path=str(path),
        meta_data={
            "parser": "txt",
            "encoding": encoding
        },
    )
def markdown_parser(path)->Document:
    text, encoding = read_text_file(path)
    return Document(
        id=md5_calculate_doc_id(str(path)),
        name=path.name,
        content=text,
        source_path=str(path),
        meta_data={
            "parser": "markdown",
            "encoding": encoding
        },
    )

parser = {".md": markdown_parser,
          ".markdown":markdown_parser,
          ".pdf":pdf_parser,
          ".jpg":picture_parser,
          ".jpeg":picture_parser,
          ".png":picture_parser,
          ".docx":docx_parser,
          ".txt":text_parser}

def md5_calculate_doc_id(path):
    return hashlib.md5(path.encode("utf-8")).hexdigest()

def clean_text(text: str) -> str:
    """统一换行符为 \\n，去掉首尾空白。"""
    text = text.replace("\r\n", "\n")  # Windows 换行
    text = text.replace("\r", "\n")     # 老式 Mac 换行
    return text.strip()

def load_directory(directory:str):
    document = []
    fail_to_parse_document = {}
    empty_document = []
    for file in Path(directory).rglob("*"):
        if file.is_file():
            suffix = file.suffix.lower()
            if parser.get(suffix):
                try:
                    temp_document = parser.get(suffix)(file)
                    temp_document.content = clean_text(temp_document.content)
                    if not temp_document.content:
                        empty_document.append(temp_document)
                    else:
                        document.append(temp_document)
                # 记录失败原因
                except Exception as e:
                    fail_to_parse_document[str(file)] = str(e)
            else:
                # 记录失败原因
                fail_to_parse_document[str(file)]=str(f"不支持{suffix}类型")
    return document,fail_to_parse_document,empty_document



if __name__ == "__main__":
    document,fail_to_parse_document,empty_document= load_directory("sample_notes")
    for j in document:
        print(j)
    print("-" * 50)
    for i in fail_to_parse_document.keys():
        print(i,":",fail_to_parse_document.get(i))
    print("-"*50)
    for k in empty_document:
        print(k)

